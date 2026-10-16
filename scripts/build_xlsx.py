"""초등_성취기준_학년별.xlsx 생성. 실행: python3 scripts/build_xlsx.py"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from norm import HERE, SUBJ_ORDER, by_grade, label, load

FONT = "맑은 고딕"
HEAD_FILL = PatternFill("solid", fgColor="2E6B62")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10)
THIN = Side(style="thin", color="D8E0DB")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
MARK = {"두 학년 공통": "FFF4D6", "두 학년 반복·심화": "EEF3F0"}

data = load()
grades = by_grade(data)
wb = Workbook()


def sheet(title, headers, widths, rows, wrap_cols=(), fills=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, Alignment(vertical="center"), GRID
    for i, row in enumerate(rows):
        ws.append(row)
        for j, c in enumerate(ws[ws.max_row], 1):
            c.font, c.border = BODY, GRID
            c.alignment = WRAP if j in wrap_cols else TOP
            if fills and fills[i]:
                c.fill = PatternFill("solid", fgColor=fills[i])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws


area = lambda r: f"({r['영역번호']}) {r['영역']}"

SEM_KEY = {"1학기": 0, "1학기 · 2학기": 1, "2학기": 2}


def sem_key(sems):
    return SEM_KEY.get(" · ".join(sems), 9)


# 학년별 시트: 학기 → 교과 → 코드 순
for g in range(1, 7):
    rows, fills = [], []
    for r, sems, units in sorted(grades[g], key=lambda x: (sem_key(x[1]), SUBJ_ORDER.index(x[0]["교과"]), x[0]["코드"])):
        rows.append([r["교과"], " · ".join(sems), "\n".join(units), area(r), r["코드"], r["성취기준"], label(r), r["근거"]])
        fills.append(MARK.get(label(r)))
    sheet(f"{g}학년", ["교과", "학기", "단원", "영역", "코드", "성취기준", "구분", "근거"],
          [8, 11, 34, 20, 13, 70, 16, 11], rows, wrap_cols=(3, 6), fills=fills)

# 과목·영역별: 교과 → 영역(교육과정 순서) → 학년 → 학기 → 코드
AREA_ORDER = {}
for r in data:  # 원본 CSV 순서대로 영역이 처음 나온 순번
    AREA_ORDER.setdefault((r["교과"], r["영역"]), len(AREA_ORDER))
by_area, fills_a, prev_area, shade = [], [], None, False
rows_a = []
for g in range(1, 7):
    for r, sems, units in grades[g]:
        rows_a.append((SUBJ_ORDER.index(r["교과"]), AREA_ORDER[(r["교과"], r["영역"])], g, sem_key(sems), r["코드"], r, sems, units))
for r in data:
    if not r["배치"]:
        rows_a.append((SUBJ_ORDER.index(r["교과"]), AREA_ORDER[(r["교과"], r["영역"])], 9, 9, r["코드"], r, [], []))
for *_, r, sems, units in sorted(rows_a, key=lambda x: x[:5]):
    key = (r["교과"], r["영역"])
    if key != prev_area:
        shade, prev_area = not shade, key
    g = next((x[2] for x in rows_a if x[5] is r and x[6] is sems), None)
    by_area.append([r["교과"], area(r), f"{g}학년" if g and g < 9 else "-", " · ".join(sems) or "-",
                    "\n".join(units) or "미배치", r["코드"], r["성취기준"], label(r), r["학년군"]])
    fills_a.append("F3F7F4" if shade else None)
sheet("과목·영역별", ["교과", "영역", "학년", "학기", "단원", "코드", "성취기준", "구분", "학년군"],
      [8, 24, 7, 11, 32, 13, 64, 16, 11], by_area, wrap_cols=(5, 7), fills=fills_a)

# 미배치
un = sorted([r for r in data if not r["배치"]], key=lambda r: (r["학년군"], SUBJ_ORDER.index(r["교과"]), r["코드"]))
sheet("미배치", ["학년군", "교과", "영역", "코드", "성취기준", "비고"], [12, 8, 22, 13, 80, 34],
      [[r["학년군"], r["교과"], area(r), r["코드"], r["성취기준"], r["구분"]] for r in un], wrap_cols=(5,))

# 원본 엑셀과 같은 순서로 대비
full = []


def full_key(r):
    if not r["배치"]:
        return (r["학년군"], 9, 9, SUBJ_ORDER.index(r["교과"]), r["코드"])
    b0 = min(r["배치"], key=lambda b: (b["학년"], b["학기"]))
    return (r["학년군"], b0["학년"], int(b0["학기"][0]), SUBJ_ORDER.index(r["교과"]), r["코드"])


for r in sorted(data, key=full_key):
    where = "\n".join(f"{b['학년']}학년 {b['학기']} {b['단원']}" for b in sorted(r["배치"], key=lambda b: (b["학년"], b["학기"]))) or "미배치"
    gs = sorted({b["학년"] for b in r["배치"]})
    full.append([r["학년군"], r["교과"], area(r), r["코드"], r["성취기준"], " · ".join(f"{x}학년" for x in gs) or "-", where, label(r)])
sheet("전체(학년군 대비)", ["학년군", "교과", "영역", "코드", "성취기준", "배치 학년", "학기·단원", "구분"],
      [12, 8, 22, 13, 64, 12, 44, 16], full, wrap_cols=(5, 7))

# 안내 + 집계(수식)
ws = wb["Sheet"]; ws.title = "안내"
lines = [
    ("초등 성취기준 학년별 배치표 (2022 개정 교육과정 · 2026학년도)", True),
    ("", False),
    ("자료", True),
    ("성취기준: 초등_성취기준_학년군별.xlsx (suseok-exam / prep-materials, 교육부 고시 제2022-33호) 중 9개 교과 523개", False),
    ("단원: 2026학년도 학교 연간 지도 계획(1~6학년 1·2학기, hwp 89개)", False),
    ("", False),
    ("배치 방법", True),
    ("성취기준 문장과 계획서의 단원·차시 내용을 대조해 학년·학기·단원을 정하고, 어긋난 항목은 직접 검토해 바로잡았습니다(근거 열).", False),
    ("자동 대조: 문장 유사도로 배치 / 검토 보정: 교사 검토로 수정 / 영역 대단원: 체육은 영역(운동·스포츠·표현) 대단원에 배치", False),
    ("", False),
    ("구분", True),
    ("단일 학년: 한 학년에서만 다룸", False),
    ("두 학년 공통: 같은 학년군의 두 학년에서 모두 다룸 (노란 칸)", False),
    ("두 학년 반복·심화: 음악·미술·체육처럼 두 학년에 걸쳐 반복·심화하는 교과 (연녹색 칸)", False),
    ("", False),
    ("빠진 자료", True),
    ("5~6학년군 과학의 '산과 염기'·'물체의 운동' 성취기준 7개는 2026학년도 계획서에 해당 단원이 없어(교육과정 전환기 편성) '미배치' 시트에 남겼습니다.", False),
    ("1~2학년 통합교과와 영어는 이번 정리에서 제외했습니다.", False),
    ("", False),
    ("시트 안내", True),
    ("과목·영역별: 교과 → 영역 → 학년 → 학기 순 / 내용 체계 도식: 교과별 영역 × 1~6학년 표 / 1~6학년: 학년별(학기 → 교과 순) / 미배치 / 전체(학년군 대비): 원본 학년군 엑셀과 대조용", False),
    ("", False),
    ("학년별 성취기준 수 (각 학년 시트에서 자동 집계)", True),
]
for t, b in lines:
    ws.append([t])
    ws.cell(ws.max_row, 1).font = Font(name=FONT, bold=b, size=13 if ws.max_row == 1 else 10)
hr = ws.max_row + 1
ws.append(["교과"] + [f"{g}학년" for g in range(1, 7)] + ["합계"])
for c in ws[hr]:
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, GRID
for s in SUBJ_ORDER:
    r = ws.max_row + 1
    ws.append([s] + [f"=COUNTIF('{g}학년'!$A:$A,$A{r})" for g in range(1, 7)] + [f"=SUM(B{r}:G{r})"])
r = ws.max_row + 1
ws.append(["합계"] + [f"=SUM({get_column_letter(c)}{hr + 1}:{get_column_letter(c)}{r - 1})" for c in range(2, 9)])
ws.append(["미배치"] + [""] * 6 + [f"=COUNTA('미배치'!$D:$D)-1"])
for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row):
    for c in row:
        c.font, c.border = Font(name=FONT, size=10, bold=row[0].value in ("합계", "미배치")), GRID
ws.cell(ws.max_row, 9, "← 2026 계획서에 해당 단원 없음").font = Font(name=FONT, size=9, color="5D6B64")
ws.column_dimensions["A"].width = 16
for c in "BCDEFGH":
    ws.column_dimensions[c].width = 9
# 내용 체계 도식: 교과별 영역 × 1~6학년 (build_matrix.py 결과 사용)
import json
mx = json.load(open(os.path.join(HERE, "matrix.json"), encoding="utf-8"))
ws = wb.create_sheet("내용 체계 도식")
ws.append(["학년별 내용 체계 도식 — 교과 › 영역 × 학년. 셀: [학기] 단원 (성취기준 코드). 사회·과학은 2022 개정 내용 영역으로 재분류"])
ws.cell(1, 1).font = Font(name=FONT, bold=True, size=12)
BAND_FILL = {1: "F3F7F5", 2: "F3F7F5", 3: "FFFFFF", 4: "FFFFFF", 5: "F6F3FA", 6: "F6F3FA"}
for S in mx["subjects"]:
    ws.append([])
    hr = ws.max_row + 1
    ws.append([f"{S['name']} ({S['count']}개)"] + [f"{g}학년" for g in range(1, 7)])
    for c in ws[hr]:
        c.font, c.fill, c.border, c.alignment = HEAD_FONT, HEAD_FILL, GRID, Alignment(vertical="center")
    for row in S["rows"]:
        n = len({s["code"] for c in row["cells"] for u in c for s in u["stds"]})
        vals = [f"{row['name']}\n({n}개)"]
        for c in row["cells"]:
            parts = []
            for u in c:
                tag = "1·2" if len(u["sems"]) > 1 or "·" in u["sems"][0] else u["sems"][0][0]
                codes = " ".join(s["code"].strip("[]") for s in u["stds"])
                parts.append(f"[{tag}] {u['u']}\n   {codes}")
            vals.append("\n".join(parts))
        ws.append(vals)
        r = ws.max_row
        lines = max(v.count("\n") + 1 for v in vals)
        ws.row_dimensions[r].height = max(30, 13 * lines)
        for j, cell in enumerate(ws[r], 1):
            cell.font, cell.border, cell.alignment = BODY, GRID, WRAP
            if j == 1:
                cell.font = Font(name=FONT, size=10, bold=True)
            elif cell.value:
                cell.fill = PatternFill("solid", fgColor=BAND_FILL[j - 1])
    if S["unplaced"]:
        ws.append(["미배치: " + ", ".join(f"{u['code']} {u['group']}" for u in S["unplaced"])])
        ws.cell(ws.max_row, 1).font = Font(name=FONT, size=9, color="B54848")
ws.column_dimensions["A"].width = 18
for c in "BCDEFG":
    ws.column_dimensions[c].width = 30
ws.freeze_panes = "B2"

wb.move_sheet("안내", offset=-(len(wb.sheetnames) - 1))
wb.move_sheet("과목·영역별", offset=-(wb.sheetnames.index("과목·영역별") - 1))
wb.move_sheet("내용 체계 도식", offset=-(wb.sheetnames.index("내용 체계 도식") - 2))

out = os.path.join(os.path.dirname(HERE), "초등_성취기준_학년별.xlsx")
wb.save(out)
print(out)
